#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.10"
# dependencies = ["pypandoc-binary>=1.13", "python-docx>=1.1"]
# ///
"""Build assets/reference.docx: pandoc's default reference restyled to match style.css.

Run again after changing the palette or fonts here. Word fonts are chosen for what
recipients have installed: Georgia (body), Aptos (structure), Consolas (code).
"""

from __future__ import annotations

import io
import subprocess
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parent.parent / "assets" / "reference.docx"

SERIF, SANS, MONO = "Georgia", "Aptos", "Consolas"
INK, INK2, MUTED = "1C1A17", "4A463F", "6E685E"
RULE, WASH, ACCENT = "E4DFD5", "F3F0E9", "23548C"


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def set_font(style, name: str) -> None:
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    for theme in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        fonts.attrib.pop(qn(theme), None)


def set_spacing_twips(style, chars: float) -> None:
    """Letter spacing in points."""
    rpr = style.element.get_or_add_rPr()
    sp = rpr.find(qn("w:spacing"))
    if sp is None:
        sp = OxmlElement("w:spacing")
        rpr.append(sp)
    sp.set(qn("w:val"), str(int(chars * 20)))


def shade(ppr_or_rpr, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr_or_rpr.append(shd)


def para_border(style, side: str, color: str, size_eighths: int, space: int = 8) -> None:
    ppr = style.element.get_or_add_pPr()
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        ppr.append(bdr)
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size_eighths))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    bdr.append(el)


def get(doc, name: str, kind=WD_STYLE_TYPE.PARAGRAPH):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, kind)


def para(doc, name, *, font=None, size=None, bold=None, italic=None, color=None,
         before=None, after=None, line=None, caps=None, align=None, keep=None, tracking=None):
    s = get(doc, name)
    if font:
        set_font(s, font)
    f = s.font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color:
        f.color.rgb = rgb(color)
        f.color.theme_color = None
    if caps is not None:
        f.all_caps = caps
    pf = s.paragraph_format
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
    if align is not None:
        pf.alignment = align
    if keep is not None:
        pf.keep_with_next = keep
    if tracking is not None:
        set_spacing_twips(s, tracking)
    return s


def table_style(doc) -> None:
    s = get(doc, "Table", WD_STYLE_TYPE.TABLE)
    set_font(s, SANS)
    s.font.size = Pt(9.5)
    el = s.element
    for old in el.findall(qn("w:tblPr")) + el.findall(qn("w:tblStylePr")):
        el.remove(old)
    tblpr = OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for side, sz, color in (("top", 0, "auto"), ("left", 0, "auto"), ("right", 0, "auto"),
                            ("insideV", 0, "auto"), ("bottom", 8, INK), ("insideH", 4, RULE)):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single" if sz else "nil")
        if sz:
            b.set(qn("w:sz"), str(sz))
            b.set(qn("w:space"), "0")
            b.set(qn("w:color"), color)
        borders.append(b)
    tblpr.append(borders)
    mar = OxmlElement("w:tblCellMar")
    for side, w in (("top", 80), ("bottom", 80), ("left", 0), ("right", 280)):
        m = OxmlElement(f"w:{side}")
        m.set(qn("w:w"), str(w))
        m.set(qn("w:type"), "dxa")
        mar.append(m)
    tblpr.append(mar)
    el.append(tblpr)

    first = OxmlElement("w:tblStylePr")
    first.set(qn("w:type"), "firstRow")
    rpr = OxmlElement("w:rPr")
    for tag, val in (("w:b", None), ("w:caps", None), ("w:color", MUTED), ("w:sz", "15"), ("w:spacing", "10")):
        e = OxmlElement(tag)
        if val:
            e.set(qn("w:val"), val)
        rpr.append(e)
    first.append(rpr)
    tcpr = OxmlElement("w:tcPr")
    tb = OxmlElement("w:tcBorders")
    bot = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "0"), ("w:color", INK)):
        bot.set(qn(k), v)
    tb.append(bot)
    tcpr.append(tb)
    first.append(tcpr)
    el.append(first)


def footer_page_number(doc) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1.1)
    section.different_first_page_header_footer = True
    p = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.font.name, run.font.size, run.font.color.rgb = SANS, Pt(8), rgb(MUTED)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            run._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            run._r.append(it)


def main() -> None:
    raw = subprocess.run(
        [pypandoc.get_pandoc_path(), "--print-default-data-file", "reference.docx"],
        check=True, capture_output=True,
    ).stdout
    doc = Document(io.BytesIO(raw))

    para(doc, "Normal", font=SERIF, size=11, color=INK, line=1.38, after=8)
    for name in ("Body Text", "First Paragraph"):
        para(doc, name, before=0, after=8)
    para(doc, "Compact", before=2, after=2)

    para(doc, "Title", font=SERIF, size=28, bold=True, color=INK, before=0, after=6, line=1.05, align=WD_ALIGN_PARAGRAPH.LEFT)
    para(doc, "Subtitle", font=SERIF, size=14, italic=True, bold=False, color=INK2, before=0, after=14, line=1.2, align=WD_ALIGN_PARAGRAPH.LEFT)
    for name in ("Author", "Date"):
        para(doc, name, font=SANS, size=8.5, color=MUTED, caps=True, tracking=0.6, before=0, after=2, align=WD_ALIGN_PARAGRAPH.LEFT)
    date = para(doc, "Date", after=26)
    para_border(date, "bottom", RULE, 4, space=14)
    para(doc, "Abstract", font=SERIF, size=11, italic=True, color=INK2, align=WD_ALIGN_PARAGRAPH.LEFT)

    para(doc, "Heading 1", font=SERIF, size=20, bold=True, color=INK, before=28, after=10, line=1.1, keep=True)
    para(doc, "Heading 2", font=SERIF, size=15, bold=True, color=INK, before=24, after=8, line=1.15, keep=True)
    para(doc, "Heading 3", font=SANS, size=12, bold=True, color=INK, before=18, after=4, line=1.2, keep=True)
    para(doc, "Heading 4", font=SANS, size=8.5, bold=True, italic=False, color=MUTED, caps=True, tracking=0.7, before=16, after=4, keep=True)
    for name in ("Heading 5", "Heading 6"):
        para(doc, name, font=SANS, size=10.5, bold=True, italic=False, color=INK, before=12, after=3, keep=True)

    bq = para(doc, "Block Text", font=SERIF, size=12, italic=True, color=INK2, before=6, after=12, line=1.35)
    bq.paragraph_format.left_indent = Inches(0.3)
    bq.paragraph_format.right_indent = Inches(0)
    para_border(bq, "left", ACCENT, 12, space=12)

    code = para(doc, "Source Code", font=MONO, size=9, color=INK, before=0, after=0, line=1.25)
    shade(code.element.get_or_add_pPr(), WASH)
    vc = get(doc, "Verbatim Char", WD_STYLE_TYPE.CHARACTER)
    set_font(vc, MONO)
    vc.font.size = Pt(9.5)
    shade(vc.element.get_or_add_rPr(), WASH)

    link = get(doc, "Hyperlink", WD_STYLE_TYPE.CHARACTER)
    link.font.color.rgb = rgb(ACCENT)
    link.font.underline = True

    para(doc, "Footnote Text", font=SERIF, size=9, color=INK2, after=4, line=1.25)
    for name in ("Image Caption", "Caption", "Table Caption"):
        para(doc, name, font=SANS, size=8.5, italic=False, color=MUTED, before=4, after=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "Captioned Figure", align=WD_ALIGN_PARAGRAPH.CENTER, keep=True)
    para(doc, "Figure", align=WD_ALIGN_PARAGRAPH.CENTER, keep=True)
    para(doc, "TOC Heading", font=SANS, size=8.5, bold=True, color=MUTED, caps=True, tracking=0.7, before=0, after=8)

    table_style(doc)
    footer_page_number(doc)
    doc.save(OUT)
    print(f"ok   {OUT}")


if __name__ == "__main__":
    main()
